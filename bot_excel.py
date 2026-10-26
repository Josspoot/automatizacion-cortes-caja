#!/usr/bin/env python3
"""Captura los cortes de caja (PDF) en un libro Excel con macros (.xlsm).

Sustituye a la plataforma web: escribe directo en las hojas INGRESOS y GASTOS
imitando a un formulario VBA del libro (inserta fila 2, copia el formato de la
fila 3, folio consecutivo I##### / G#####, guarda el libro). Usa Excel de
verdad via AppleScript para no romper macros, tablas ni segmentadores.

Pensado para un libro con la misma estructura de hojas/columnas que el tuyo;
ajusta las referencias de celda si tu plantilla difiere.

Uso:
    python bot_excel.py "cortes_pdf/CC 02-07-2026.pdf"            # ENSAYO
    python bot_excel.py "cortes_pdf/CC 02-07-2026.pdf" --real     # escribe
    python bot_excel.py "... .pdf" --real --solo 5                # solo fila 5

Convencion de fechas (2026-07-14): la FECHA registrada es la MISMA del corte.
"""
import argparse
import json
import re
import subprocess
import sys
import unicodedata
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import leer_corte
import config

BASE = Path(__file__).parent
XLSM = config.XLSM_PATH
CEREBRO = config.CONOCIMIENTO_EXCEL_PATH
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def _norm(t):
    t = unicodedata.normalize("NFD", t or "")
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().upper()


def fecha_operacion(fecha_pdf):
    """'DD/MM/YYYY' del corte -> date. Se registra con la MISMA fecha del PDF
    (decidido 2026-07-14; la convencion 'dia anterior' quedo descartada)."""
    d, m, y = fecha_pdf.strip().split("/")
    return date(int(y), int(m), int(d))


def serial_excel(d):
    """date -> numero de serie de Excel (la celda ya tiene formato de fecha)."""
    return (d - date(1899, 12, 30)).days


def cargar_catalogos():
    """Lee del .xlsm los proveedores (NC -> datos), clientes, folios siguientes
    y las fechas que ya tienen filas (para no capturar un dia dos veces)."""
    import warnings
    from datetime import datetime
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(XLSM, read_only=True)
    provs = {}
    for row in wb["PROVEEDORES"].iter_rows(min_row=2, max_col=10, values_only=True):
        if row[2]:
            provs[_norm(str(row[2]))] = {
                "rs": str(row[0] or "-").strip(), "rfc": str(row[1] or "-").strip(),
                "nc": str(row[2]).strip(), "sub": str(row[8] or "-").strip(),
                "tg": str(row[9] or "-").strip()}
    clientes = set()
    for row in wb["CLIENTES"].iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            clientes.add(_norm(str(row[0])))
    cnt, fechas = {}, {}
    for hoja in ("INGRESOS", "GASTOS"):
        n, fs = 0, set()
        for row in wb[hoja].iter_rows(max_col=16, values_only=True):
            if row[0] not in (None, ""):
                n += 1
            f = row[15]
            if isinstance(f, datetime):
                fs.add(f.date())
        cnt[hoja] = n
        fechas[hoja] = fs
    wb.close()
    return provs, clientes, cnt, fechas


def mapear(mov, cer, provs, clientes):
    """Movimiento del corte -> (hoja, valores A..Q) o (None, motivo)."""
    seg = _norm(mov.segmento)
    f_op = fecha_operacion(mov.fecha)
    ser = serial_excel(f_op)
    mes = MESES[f_op.month - 1]
    efectivo = "EFECT" in _norm(mov.forma_pago or mov.banco or "EFECTIVO")
    forma = "EFECTIVO" if efectivo else "TRANSFERENCIA"
    codigo = 1 if efectivo else 2
    cuenta = "-" if efectivo else "AZTECA"
    concepto = re.sub(r"\s+", " ", mov.concepto).strip().upper()

    if mov.tipo in ("INGRESO", "DEPOSITO"):
        info = cer["ingresos"].get(seg)
        if not info:
            return None, f"sin mapeo de INGRESO para '{mov.segmento}'"
        if _norm(info["cliente"]) not in clientes:
            return None, f"cliente '{info['cliente']}' no existe en hoja CLIENTES"
        vals = [None, info["cliente"], "-", "-", info["tipo"], ser, ser,
                "-", "-", "-", concepto, forma, codigo, cuenta,
                mov.importe, ser, mes]
        return ("INGRESOS", vals), None

    # GASTO. La nomina detectada (variables tipo DISPERSION/REEMBOLSOS nomina)
    # tiene prioridad sobre el mapeo del segmento.
    info, nc = None, None
    if "NOMINA" in _norm(mov.contacto or ""):
        nc = cer["nomina_web_a_nc"].get(_norm(mov.contacto))
    regla = None
    if not nc:
        info = cer["gastos"].get(seg)
        if info:
            # reglas por contenido: pueden cambiar proveedor y/o subconcepto
            regla = next((r for r in info.get("reglas_contenido", [])
                          if any(p in _norm(concepto) for p in r["si"])), None)
            if info.get("solo_con_regla") and regla is None:
                return None, f"'{mov.segmento}' es variable y ninguna regla aplica"
            nc = (regla or {}).get("nc") or info.get("nc")
    if not nc:
        return None, f"sin mapeo de GASTO para '{mov.segmento}'"
    prov = provs.get(_norm(nc))
    if not prov:
        return None, f"proveedor '{nc}' no existe en hoja PROVEEDORES"
    sub, tg = prov["sub"], prov["tg"]
    if info:
        sub, tg = info.get("sub", sub), info.get("tg", tg)
    if regla:
        sub, tg = regla.get("sub", sub), regla.get("tg", tg)
    vals = [None, prov["rs"], prov["rfc"], prov["nc"], concepto, sub,
            tg, concepto, forma, codigo, cuenta, "-",
            mov.importe, "CONTADO", "PAGADO", ser, mes]
    return ("GASTOS", vals), None


def _as_str(v):
    if isinstance(v, str):
        return '"' + v.replace("\\", "\\\\").replace('"', '\\"') + '"'
    return str(v)


def applescript_preparar():
    """Verifica que el libro abierto sea EL DE DOCUMENTOS y lo guarda; si no
    esta abierto, lo abre. Devuelve error si hay otro con el mismo nombre."""
    return f'''
    tell application "Microsoft Excel"
        set rutaOK to "{XLSM}"
        if exists workbook "{XLSM.name}" then
            set rutaReal to (full name of workbook "{XLSM.name}") as text
            if rutaReal is not equal to rutaOK then
                error "ABIERTO_OTRO:" & rutaReal
            end if
            save workbook "{XLSM.name}"
        else
            open workbook workbook file name rutaOK
            delay 1
        end if
        return "OK"
    end tell
    '''


def applescript_escribir(filas):
    """filas = [(hoja, [17 valores A..Q]), ...] en orden del PDF."""
    cuerpo = []
    for hoja, vals in filas:
        sets = "\n".join(
            f'            set value of range "{chr(65+i)}2" of ws to {_as_str(v)}'
            for i, v in enumerate(vals))
        cuerpo.append(f'''
        set ws to worksheet "{hoja}" of wb
        insert into range (entire row of range "A2" of ws)
        copy range (range "A3:Q3" of ws)
        paste special (range "A2:Q2" of ws) what paste formats
{sets}''')
    return f'''
    tell application "Microsoft Excel"
        set screen updating to false
        set wb to workbook "{XLSM.name}"
{"".join(cuerpo)}
        set screen updating to true
        save wb
        return "OK"
    end tell
    '''


def correr_osascript(script):
    r = subprocess.run(["osascript", "-e", script],
                       capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip())
    return r.stdout.strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdf")
    ap.add_argument("--real", action="store_true", help="escribir en el Excel")
    ap.add_argument("--solo", type=int, help="capturar SOLO la fila N del PDF")
    ap.add_argument("--forzar", action="store_true",
                    help="escribir aunque ese dia ya tenga filas en el Excel")
    args = ap.parse_args()

    cer = json.loads(CEREBRO.read_text())
    movs, fecha_doc, total_inicio = leer_corte.leer_pdf(args.pdf)
    for m in movs:
        leer_corte.aplicar_reglas(m)

    cola = [m for m in movs if m.accion in ("CAPTURAR", "PENDIENTE")
            and m.importe > 0]
    if args.solo:
        cola = [m for m in cola if int(m.n.strip("/ ")) == args.solo]

    if args.real:
        print("Preparando Excel (verifica que sea el libro de Documentos)...")
        try:
            correr_osascript(applescript_preparar())
        except RuntimeError as e:
            if "ABIERTO_OTRO" in str(e):
                print("✗ Tienes abierto OTRO archivo con el mismo nombre "
                      "(el de WhatsApp). Cierralo y abre el de Documentos, "
                      "o dejalo cerrado y el bot lo abre solo.")
                print("  " + str(e))
                sys.exit(1)
            raise

    provs, clientes, cnt, fechas = cargar_catalogos()
    n_ing, n_gas = cnt["INGRESOS"], cnt["GASTOS"]

    f_op = fecha_operacion(fecha_doc)
    ya_estaba = f_op in (fechas["INGRESOS"] | fechas["GASTOS"])
    if args.real and ya_estaba and not (args.forzar or args.solo):
        print(f"⚠ El {f_op:%d/%m/%Y} YA tiene filas en el Excel: no escribo nada "
              f"para no duplicar. (--forzar para escribir de todos modos, "
              f"--solo N para recapturar una fila)")
        sys.exit(3)
    print(f"\nCorte {fecha_doc}  ->  FECHA registrada {f_op:%d/%m/%Y} ({MESES[f_op.month-1]})")
    if total_inicio:
        print(f"TOTAL INICIO DIA (cuadre): ${total_inicio:,.2f}")
    print(f"Folios siguientes: I{n_ing:05d} / G{n_gas:05d}\n")

    # El cerebro de EXCEL manda: si mapea, se captura; si no, queda manual.
    listas, pendientes = [], []
    for m in cola:
        res, motivo = mapear(m, cer, provs, clientes)
        if res is None:
            pendientes.append((m, motivo or m.nota or "pendiente", None))
            continue
        listas.append((m, res))

    folio = {"INGRESOS": n_ing, "GASTOS": n_gas}
    filas_final = []
    print(f"{'fila':>4} {'folio':7} {'hoja':8} {'contraparte':32} {'sub/tipo':34} {'pago':13} {'monto':>12}")
    for m, (hoja, vals) in listas:
        pref = "I" if hoja == "INGRESOS" else "G"
        vals[0] = f"{pref}{folio[hoja]:05d}"
        folio[hoja] += 1
        filas_final.append((hoja, vals))
        contraparte = vals[1] if hoja == "GASTOS" else vals[1]
        subt = vals[5] if hoja == "GASTOS" else vals[4]
        print(f"{m.n.strip('/'):>4} {vals[0]:7} {hoja:8} {str(contraparte)[:32]:32} "
              f"{str(subt)[:34]:34} {vals[8] if hoja=='GASTOS' else vals[11]:13} "
              f"${m.importe:>10,.2f}")

    if pendientes:
        print(f"\n-- NO se capturan ({len(pendientes)}) --")
        for m, razon, _ in pendientes:
            print(f"  fila {m.n.strip('/'):>3} ${m.importe:>10,.2f}  {m.concepto[:52]:52} -> {razon}")

    if not args.real:
        print(f"\nENSAYO: se escribirian {len(filas_final)} filas. "
              f"Corre con --real para escribir.")
        return

    if not filas_final:
        print("\nNada que escribir.")
        return

    print(f"\nEscribiendo {len(filas_final)} filas en el Excel...")
    correr_osascript(applescript_escribir(filas_final))
    print(f"✔ {len(filas_final)} filas escritas y libro guardado.")
    if pendientes:
        print(f"⚠ Quedaron {len(pendientes)} pendientes (arriba) para captura manual.")


if __name__ == "__main__":
    main()
