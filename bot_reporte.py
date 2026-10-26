#!/usr/bin/env python3
"""Captura los REPORTES DE SUCURSAL (xlsx "INGRESOS DEL ..."/"GASTOS DEL ...")
en las DOS plataformas: el Excel con macros y el ERP web (ver config.py / .env).

A diferencia de los cortes PDF, estos reportes traen el detalle POR TICKET
(folio, productos, forma de pago). Reglas acordadas (2026-07-18):
  - FOLIO INTERNO en el Excel: servicios de taller (CESEM) llevan su "SER-000"
    de las observaciones; el resto lleva el numero de ticket. El folio va
    TAMBIEN dentro del concepto/detalle.
  - Pagos mixtos (efectivo + banco) se dividen en una fila por forma de pago.
  - Tarjeta debito -> TERMINAL, tarjeta credito -> TARJETA CREDITO (codigo 2,
    cuenta AZTECA) en el Excel; en la web se resuelve la opcion del combo al
    vuelo la primera vez.
  - GASTOS: solo se capturan los de fechas que aun NO tienen filas en el libro
    (los de dias con corte PDF ya estan capturados). Ingresos se deduplican
    por numero de ticket en el detalle.

Uso:
    python bot_reporte.py                      # ENSAYO con los xlsx mas nuevos
    python bot_reporte.py --excel              # escribe en el Excel
    python bot_reporte.py --web                # captura en la web
    python bot_reporte.py --web --desde 12     # reanudar web desde el mov 12
    python bot_reporte.py --web --solo 7       # solo el movimiento 7
"""
import argparse
import functools
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import bot_excel
import config
from leer_corte import Movimiento

print = functools.partial(print, flush=True)

BASE = Path(__file__).parent
DOCS = Path.home() / "Documents"
# registro por plataforma: que items del reporte ya se capturaron y donde
# (el candado del Excel no sirve para la web: en cuanto escribimos el libro,
#  todo "ya esta en el Excel" aunque la web siga sin capturarse)
REGISTRO = config.REGISTRO_REPORTE_PATH
CONOC = json.loads(config.CONOCIMIENTO_PATH.read_text(encoding="utf-8"))
CER_X = json.loads(config.CONOCIMIENTO_EXCEL_PATH.read_text(encoding="utf-8"))
MESES = bot_excel.MESES

# formas de pago del reporte: (columna, banco, forma_excel, codigo, cuenta, forma_web)
# la forma_web "TARJETA*" es un marcador: se resuelve contra el combo real.
FORMAS = [
    (11, "EFECTIVO", "EFECTIVO", 1, "-", "Efectivo"),
    (13, "AZTECA", "TERMINAL", 2, "AZTECA", "TARJETA_DEBITO"),
    (14, "AZTECA", "TARJETA CREDITO", 2, "AZTECA", "TARJETA_CREDITO"),
    (15, "AZTECA", "TRANSFERENCIA", 2, "AZTECA", "Transferencia"),
]


def _norm(t):
    t = unicodedata.normalize("NFD", str(t or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().upper()


def _fecha(v):
    if isinstance(v, datetime):
        return v.date()
    return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()


def _limpiar_producto(p):
    """'FOCO LED ... 1 unidades, $100 con 0% de desc.' -> 'FOCO LED ...'"""
    p = re.sub(r"\s*\d+\s+unidades.*$", "", str(p)).strip(" ,")
    return re.sub(r"\s+", " ", p)[:45].strip()


# ---------------------------------------------------------------------------
# LECTURA DE LOS REPORTES
# ---------------------------------------------------------------------------
def leer_ingresos(ruta):
    """Tickets del reporte de ventas -> lista de dicts (uno por ticket)."""
    import openpyxl
    ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    tickets = []
    for i, r in enumerate(filas):
        if not str(r[0] or "").isdigit():
            continue
        prods, obs = [str(r[3] or "")], str(r[4] or "")
        j = i + 1
        while j < len(filas) and not str(filas[j][0] or "").isdigit():
            if filas[j][3]:
                prods.append(str(filas[j][3]))
            if filas[j][4]:
                obs += " " + str(filas[j][4])
            j += 1
        pagos = []
        for col, banco, fx, cod, cta, fw in FORMAS:
            monto = float(r[col] or 0)
            if monto > 0:
                pagos.append((monto, banco, fx, cod, cta, fw))
        tickets.append({
            "folio": int(r[0]), "fecha": _fecha(r[6]), "prods": prods,
            "obs": re.sub(r"\s+", " ", obs).strip(), "total": float(r[21] or 0),
            "pagos": pagos})
    return tickets


def _hoja_efectivo(wb):
    """Reportes viejos: 1 sola hoja (gastos en efectivo). Reportes nuevos
    (desde 2026-07-21): 2 hojas, EFECTIVO y BANCO -> se busca por nombre."""
    for name in wb.sheetnames:
        if "EFECTIVO" in name.upper():
            return wb[name]
    return wb.worksheets[0]


def leer_gastos(ruta):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = _hoja_efectivo(wb)
    gastos = []
    for r in ws.iter_rows(values_only=True):
        if not str(r[0] or "").startswith("G-"):
            continue
        gastos.append({
            "folio": str(r[0]), "fecha": _fecha(r[3]),
            "detalle": re.sub(r"\s+", " ", str(r[8] or "")).strip().upper(),
            "tipo_reporte": str(r[2] or ""), "monto": abs(float(r[7] or 0)),
            "fuente": "EFECTIVO"})
    return sorted(gastos, key=lambda g: (g["fecha"], g["folio"]))


# tipos de movimiento de banco que son traspasos internos caja<->banco (no
# son gasto ni ingreso real: el dinero solo cambia de "bolsillo")
TRASPASOS_INTERNOS_BANCO = ("TRANSFERENCIA A BANCO", "TRANSFERENCIA A EFECTIVO")


def leer_banco(ruta):
    """Hoja BANCO (reportes nuevos): estado de cuenta con TODOS los
    movimientos. Solo interesan los EGRESOS reales (comisiones, pagos a
    proveedores, renta...); los 'Ingreso' son depositos que confirman ventas
    con tarjeta/transferencia YA capturadas por ticket en INGRESOS, y los
    traspasos internos caja<->banco no se registran."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if "BANCO" not in wb.sheetnames:
        return []
    rows = list(wb["BANCO"].iter_rows(values_only=True))
    egresos = []
    for i, r in enumerate(rows):
        if not r[0] or not str(r[0]).isdigit():
            continue
        clasif = rows[i + 1][2] if i + 1 < len(rows) else ""
        movimiento = str(r[2] or "").upper()
        if clasif != "Egreso" or any(t in movimiento for t in TRASPASOS_INTERNOS_BANCO):
            continue
        egresos.append({
            "folio": f"B-{r[0]}", "fecha": _fecha(r[4]),
            "detalle": re.sub(r"\s+", " ", str(r[10] or "")).strip().upper(),
            "tipo_reporte": str(r[11] or ""), "monto": abs(float(r[3] or 0)),
            "fuente": "BANCO"})
    return sorted(egresos, key=lambda g: (g["fecha"], g["folio"]))


# ---------------------------------------------------------------------------
# MAPEO
# ---------------------------------------------------------------------------
def clasificar_ticket(t):
    """-> ('CESEM'|'REFACCIONARIA'|'VENTA MOTOCICLETAS', concepto, folio_interno)"""
    folio, obs = t["folio"], t["obs"]
    p0 = _norm(t["prods"][0])
    sers = re.findall(r"\bSER[\s\-]*(\d+)", obs, re.I)
    ser = "/".join(f"SER-{n}" for n in sers)
    if p0.startswith("SERVICIOS DE TALLER"):
        desc = _norm(re.sub(r"\bSER[\s\-]*\d+", "", obs, flags=re.I)).strip("- ")
        desc = desc or "SERVICIO DE TALLER"
        conc = f"CESEM - {desc} - {folio}" + (f" {ser}" if ser else "")
        return "CESEM", conc, (ser or str(folio))
    # el nombre del producto no siempre trae "Motocicleta" al inicio (ej.
    # "280Z GRIS CON GPS ORIGEN"): SERIE:/CHASIS: en observaciones es la
    # señal confiable de que es una venta de moto.
    es_moto = p0.startswith("MOTOCICLETA") or bool(re.search(r"\b(SERIE|CHASIS)\s*:", obs, re.I))
    if es_moto:
        detalle = _norm(obs) or _limpiar_producto(t["prods"][0]).upper()
        return "VENTA MOTOCICLETAS", f"VENTA MOTOCICLETAS - {detalle} - {folio}", str(folio)
    prods = " + ".join(_limpiar_producto(p).upper() for p in t["prods"] if _limpiar_producto(p))
    if obs:
        prods += f" ({_norm(obs)})"
    return "REFACCIONARIA", f"REFACCIONARIA - {prods} - {folio}", str(folio)


# gastos: (patrones en el detalle) -> (nc excel, clave web en conocimiento.json
# o tupla explicita (proveedor,cat1,cat2,cat3) o None = manual en web)
GASTOS_MAPA = [
    (["TLAPALERIA SILICON", "TLAPALERIA IMPERMEABIL", "SILICON PARA LAMINAS",
      "SILICON PARA TECHO"], "TLAPALERIA", "TLAPALERIA TECHO"),
    (["TLAPALERIA"], "TLAPALERIA", "TLAPALERIA"),
    (["TOTAL PLAY"], "TOTAL PLAY", "TOTAL PLAY"),
    (["HORAS EXTRAS", "VIATICOS"], "VIATICOS",
     ("Otros", "Gastos Variables de Operacion", "Viaticos", None)),
    (["NOMINA PRODUCTIVA"], "NOMINA PRODUCTIVA", "NOMINA PRODUCTIVA"),
    (["NOMINA COMERCIAL", "3% DE NOMINA", "NOMINA MARIANA"], "NOMINA COMERCIAL", "NOMINA COMERCIAL"),
    (["NOMINA ADMIN"], "NOMINA ADMINISTRATIVA", "NOMINA ADMINISTRATIVA"),
    (["GASOLINA"], "GASOLINA", "GASOLINA"),
    (["GABRIEL H"], "GABRIEL HUCHIM MANTENIMIENTOS", "GABRIEL PLOMERIA"),
    (["PAPEL DE BA"], "PAC", "PAC"),                       # limpieza e higiene
    (["FIBRA PARA TRASTES", "FIBRA DE TRASTES", "VASOS DESECHABLES"], "PAC", "PAC"),
    (["GALLETAS", "COMIDA", "CAFE", "COFFE", "ROSCA", "CAZOLA"], "COMIDA DE PERSONAL", "COMIDA PERSONAL"),
    (["AGUAS PARA REUNION", "AGUAS REUNION", "GARRAFONES DE AGUA", "GARRAFON"], "COMIDA DE PERSONAL", "COMIDA PERSONAL"),
    (["AISA", "PAPELERIA"], "PAPELERIAS", "AISA"),
    (["SAMS"], "PAC", "PAC"),
    (["SELEM"], "REFACCIONARIA", "SELEM"),
    (["PLAZA"], "REFACCIONARIA", "PLAZA DE LAS MOTOS"),
    (["COPYJET"], "COPYECO", "COPYJET IMPRESORAS"),
    (["RENTA"], "RENTA DE LOCALES", "RENTA DE LOCALES"),
    (["TELCEL"], "TELCEL", "TELCEL"),
    (["RODRIGO MATOS"], "RODRIGO MATOS", "RODRIGO MATOS"),
    (["ITALIKA MEXICO", "COMERCIALIZADORA"], "ITALIKA MEXICO", "ITALIKA MEXICO"),
    (["NET PAY", "NETPAY"], "COMISIONES BANCARIAS", "NETPAY"),
    (["CLARA - CANVA", "CLARA -CANVA"], "SOFTWARE", "SOFTWARE CANVA"),
    (["CLARA - GOOGLE"], "SOFTWARE", "SOFTWARE GOOGLE"),
    (["CLARA- CLAUDE", "CLARA - CLAUDE"], "SOFTWARE", "SOFTWARE GENERICO"),
    (["CLARA - GAMMA"], "SOFTWARE", "SOFTWARE GAMMA"),
    (["EAYSWHATS", "EASYWHA"], "SOFTWARE", "SOFTWARE EASYWHA"),
    (["CLARA - GASOLINA"], "GASOLINA", "GASOLINA ARMADOS"),
    (["CLARA - SOLUCIONES SOFTWARE"], "SOFTWARE", "SOFTWARE GENERICO"),
    (["FACEBOOK - CLARA", "CLARA - PUBLICIDAD"], "FACEBOOK ADS", "FACEBOOK ADS"),
    (["MENSUALIDAD CLARA"], "CLARA TDC", "TC CLARA"),
    (["CLARA", "TDC"], "CLARA TDC", "TC CLARA"),
    (["REEMBOLSO MS"], "REEMBOLSO", "REEMBOLSO MS"),
    (["SERVI LIMPIA", "SERVI-LIMPIA", "SERVILIMPIA"], "SERVILIMPIA", "SERVI-LIMPIA"),
    (["PRESTAMO DE NOMINA", "PRESTAMOS DE NOMINA"], "PRESTAMOS DE NOMINA", "PRESTAMO"),
    (["IMSS", "INFONAVIT"], "IMMS", "IMSS"),
    (["SEGURO BANORTE", "SEGURO BANCO", "POLIZA SEGURO"], "POLIZA SEGUROS", "SEGURO BANCO"),
    (["ALLY REDES"], "ALLY REDES SOCIALES ADD MANAGER", "PUBLICIDAD"),
    (["PROVISIONES AGUINALDOS", "AGUINALDOS"], "NOMINA ADMINISTRATIVA",
     ("Nomina Administrativa", "Gastos Fijos de Operacion - Personal",
      "Nominas- Sueldos y salarios", "Nomina Administrativa")),
    (["USO BANCA"], "COMISIONES BANCARIAS", "NETPAY"),
    (["RODRIGO EFECTIVO"], "VIATICOS",
     ("Otros", "Gastos Variables de Operacion", "Viaticos", None)),
    (["APYMSA", "EL PEDAL LOCO"], "REFACCIONARIA", "SELEM"),
    (["PANELES SOLARES"], "SOLAR INSURANCE", "PANELES SOLARES"),
    (["TELMEX"], "TELMEX", "TELMEX"),
    (["MERCADO LIBRE"], "MERCADO LIBRE", "MERCADO LIBRE"),
    (["PLATZI"], "PLATZI", "PLATZI"),
    (["PROGRAMADOR"], "PROGRAMADOR", "PROGRAMADOR"),
    (["OTROS TRABAJO ARMADO", "OTROS TRABAJO ACTIVADO"], "OTROS TRABAJOS", "OTROS TRABAJOS ACTIVADOS"),
    (["EDER"], "OTROS TRABAJOS", "OTROS TRABAJOS ACTIVADOS"),
    (["USO TERMINAL"], "COMISIONES BANCARIAS", "NETPAY"),
    (["JAPAY"], "JAPAY", "JAPAY"),
    (["BOLSAS DE BASURA", "BOLSAS PARA BASURA"], "PAC", "PAC"),
    (["COMISIONES CREDITOS", "CREDITOS BAZ"], "COMISIONISTAS", "COMISIONISTAS CREDITOS"),
    (["COMBUSTIBLES ARMADO", "COMBUSTIBLES ACTIVADO", "GASOLINA ARMADO", "GASOLINA ACTIVADO",
      "GAS ACTIVADO"], "GASOLINA", "GASOLINA ARMADOS"),
    (["DILIGENCIAS"], "GASOLINA", "GASOLINA"),
    (["GASTOS CORPORATIVOS"], "GASTOS CORPORATIVOS", "ACTIVADOS"),
    (["TECNOMUNDO"], "TECNOMUNDO", "TECNOMUNDO"),
    (["BEPENSA"], "BEPENSA", "BEPENSA AGUA"),
    (["PAQUETERIA"], "PAQUETERIA", "PAQUETERIA"),
    (["CERTIFICACION"], "CERTIFICACIONES", "CERTIFICACIONES"),
    (["UBER"], "UBER", "UBER"),
    (["LIQUIDACION"], "LIQUIDACION", "LIQUIDACION"),
    (["PROMOGRAPHIC"], "PROMOGRAPHIC", "PROMOGRAPHIC"),
    (["RESERVAS AGUINALDO", "RESERVA AGUINALDO"], "RESERVAS AGUINALDO", "RESERVAS AGUINALDO"),
    (["RESERVAS"], "RESERVAS", "RESERVAS"),
    (["STEREN"], "STEREN", "STEREN"),
    (["RED CESIT"], "RED CESIT", "RED CESIT GASTO"),
    (["ARMANDO"], "ARMANDO", "ARMANDO"),
    (["CFE"], "CFE", "CFE"),
    (["COMPUTEX"], "COMPUTEX", "COMPUTEX"),
    (["FACEBOOK"], "FACEBOOK ADS", "FACEBOOK ADS"),
    (["SOLUCIONES SOFTWARE", "SOLUCIONES DBC"], "SOFTWARE", "SOFTWARE GENERICO"),
    (["PARA ACTIVADOS", "LTS PARA ACTIVADOS", "LTS ACTIVADO"], "GASOLINA", "GASOLINA ARMADOS"),
    (["COMISIONES BANCARIAS"], "COMISIONES BANCARIAS", "NETPAY"),
    (["COMISIONISTAS"], "COMISIONISTAS", "COMISIONISTAS"),
    (["COMEX"], "COMEX", "COMEX"),
    (["ESTACIONAMIENTO"], "OTROS TRABAJOS", "OTROS TRABAJOS ACTIVADOS"),
    (["MONITOREO Y ALARMAS", "911- MENSUALIDAD ALARMAS", "911 ALARMAS"], "911 ALARMAS", "911 ALARMAS"),
]


def mapear_gasto(detalle):
    for patrones, nc, web in GASTOS_MAPA:
        if any(p in _norm(detalle) for p in patrones):
            if isinstance(web, str):
                e = CONOC["gastos"][web]
                web = (e.get("proveedor"), e.get("cat1"), e.get("cat2"), e.get("cat3"))
                if nc == "NOMINA COMERCIAL" and not web[3]:
                    web = (web[0], web[1], web[2], "Nomina Comercial")
            return nc, web
    return None, None


# ---------------------------------------------------------------------------
# CONSTRUCCION DE MOVIMIENTOS
# ---------------------------------------------------------------------------
def construir(ing_path, gas_path, provs, detalles_prev, fechas_gastos_prev):
    """-> lista de dicts {mov (web), vals (excel), hoja, omitido}"""
    items = []
    for t in leer_ingresos(ing_path):
        seg, conc, folio_int = clasificar_ticket(t)
        omit = None
        if any(re.search(rf"\b{t['folio']}\b", d) for d in detalles_prev["INGRESOS"]):
            omit = f"ticket {t['folio']} ya esta en el Excel"
        if abs(sum(p[0] for p in t["pagos"]) - t["total"]) > 0.01:
            omit = f"pagos no cuadran con el total (${t['total']:,.2f})"
        info_x = CER_X["ingresos"][seg]
        info_w = CONOC["ingresos"][seg]
        ser = bot_excel.serial_excel(t["fecha"])
        for monto, banco, fx, cod, cta, fw in t["pagos"]:
            mov = Movimiento(
                n="?", fecha=t["fecha"].strftime("%d/%m/%Y"), banco=banco,
                concepto=conc, sucursal="DOLORES OTERO", importe=monto,
                es_ingreso=True, caja="001" if banco == "EFECTIVO" else "002")
            mov.tipo, mov.accion = "INGRESO", "CAPTURAR"
            mov.contacto = info_w["cliente"]
            mov.cat1, mov.cat2, mov.cat3 = (info_w.get("cat1") or "",
                                            info_w.get("cat2") or "",
                                            info_w.get("cat3") or "")
            mov.forma_pago = fw
            mov.cuenta = "Banco Azteca" if cta == "AZTECA" else ""
            vals = [None, info_x["cliente"], "-", "-", info_x["tipo"], ser, ser,
                    "-", folio_int, "-", conc, fx, cod, cta, monto, ser,
                    MESES[t["fecha"].month - 1]]
            items.append({"mov": mov, "hoja": "INGRESOS", "vals": vals,
                          "fecha": t["fecha"], "orden": t["folio"], "omitido": omit})

    for g in leer_gastos(gas_path) + leer_banco(gas_path):
        de_banco = g["fuente"] == "BANCO"
        omit = None
        if g["fecha"] in fechas_gastos_prev:
            omit = f"el {g['fecha']:%d/%m} ya tiene gastos en el Excel (corte PDF)"
        nc, web = mapear_gasto(g["detalle"])
        conc = g["detalle"]
        mov = Movimiento(
            n="?", fecha=g["fecha"].strftime("%d/%m/%Y"),
            banco="AZTECA" if de_banco else "EFECTIVO",
            concepto=conc, sucursal="DOLORES OTERO", importe=g["monto"],
            es_ingreso=False, caja="002" if de_banco else "001")
        mov.tipo = "GASTO"
        if web:
            mov.accion = "CAPTURAR"
            mov.contacto, mov.cat1, mov.cat2, mov.cat3 = [w or "" for w in web]
        else:
            mov.accion, mov.nota = "PENDIENTE", "proveedor/categoria web por definir"
        mov.forma_pago = "Transferencia" if de_banco else "Efectivo"
        mov.cuenta = "Banco Azteca" if de_banco else ""
        ser = bot_excel.serial_excel(g["fecha"])
        vals = None
        if nc and not omit:
            prov = provs.get(_norm(nc))
            if prov:
                forma_x, cod_x, cta_x = (("TRANSFERENCIA", 2, "AZTECA") if de_banco
                                         else ("EFECTIVO", 1, "-"))
                vals = [None, prov["rs"], prov["rfc"], prov["nc"], conc,
                        prov["sub"], prov["tg"], conc, forma_x, cod_x, cta_x, "-",
                        g["monto"], "CONTADO", "PAGADO", ser,
                        MESES[g["fecha"].month - 1]]
            else:
                omit = f"proveedor '{nc}' no existe en PROVEEDORES"
        elif not nc and not omit:
            omit = "sin mapeo de gasto para el Excel"
        items.append({"mov": mov, "hoja": "GASTOS", "vals": vals,
                      "fecha": g["fecha"], "orden": g["folio"], "omitido": omit})

    items.sort(key=lambda x: (x["fecha"], str(x["orden"])))
    for i, it in enumerate(items, 1):
        it["mov"].n = str(i)
    return items


def clave_item(it):
    m = it["mov"]
    return f"{it['fecha']:%Y-%m-%d}|{it['orden']}|{m.banco}|{m.importe:.2f}"


def cargar_registro():
    try:
        return json.loads(REGISTRO.read_text(encoding="utf-8"))
    except Exception:
        return {"excel": {}, "web": {}}


def guardar_registro(reg):
    REGISTRO.write_text(json.dumps(reg, indent=1, ensure_ascii=False),
                        encoding="utf-8")


def detalles_existentes():
    """Detalles ya escritos en el libro (para no duplicar tickets)."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(bot_excel.XLSM, read_only=True)
    out = {}
    for hoja, col in (("INGRESOS", 10), ("GASTOS", 7)):
        vals = set()
        for row in wb[hoja].iter_rows(min_row=2, max_col=col + 1, values_only=True):
            if row[col]:
                vals.add(str(row[col]))
        out[hoja] = vals
    wb.close()
    return out


# ---------------------------------------------------------------------------
# WEB
# ---------------------------------------------------------------------------
def resolver_forma_tarjeta(page, dlg, credito, cache={}):
    """Averigua (una vez) como se llama la opcion de tarjeta en el combo web."""
    clave = "credito" if credito else "debito"
    if clave in cache:
        return cache[clave]
    import bot as bot_web
    dlg.locator('label[for="formaPago"] + button[role="combobox"]').click()
    page.wait_for_timeout(500)
    ops = page.get_by_role("option")
    textos = [bot_web.limpiar_texto(ops.nth(i).inner_text() or "")
              for i in range(ops.count())]
    print(f"    (opciones de forma de pago: {textos})")
    prefer = (["TARJETA CREDITO", "TARJETA DE CREDITO", "CREDITO", "TARJETA", "TERMINAL"]
              if credito else
              ["TARJETA DEBITO", "TARJETA DE DEBITO", "DEBITO", "TARJETA", "TERMINAL", "TPV"])
    elegido = None
    for cand in prefer:
        for tx in textos:
            if cand in _norm(tx):
                elegido = tx
                break
        if elegido:
            break
    if elegido:
        ops.filter(has_text=elegido).first.click()
        page.wait_for_timeout(300)
    else:
        dlg.locator('label[for="formaPago"] + button[role="combobox"]').click()
        page.wait_for_timeout(300)
    cache[clave] = elegido
    return elegido


def capturar_web(items, reg, desde=1, solo=None):
    import bot as bot_web
    from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

    cola = [it for it in items if not it["omitido_web"]]
    if solo:
        cola = [it for it in cola if int(it["mov"].n) == solo]
    else:
        cola = [it for it in cola if int(it["mov"].n) >= desde]
    print(f"\n  WEB: {len(cola)} movimientos por capturar")

    with sync_playwright() as p:
        ctx = bot_web.abrir_navegador(p)
        if ctx is None:
            return
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        if "movimientos" not in page.url:
            page.goto(bot_web.URL_MOV, wait_until="domcontentloaded")
            page.wait_for_timeout(2500)
        if "auth/login" in page.url:
            print("  (!) Inicia sesion en la ventana (espero 60s)...")
            page.wait_for_timeout(60000)

        ok = err = 0
        for i, it in enumerate(cola, 1):
            m = it["mov"]
            print(f"\n[{i}/{len(cola)}]  mov {m.n}  {m.tipo}  {m.forma_pago:16} "
                  f"{m.contacto or '(tu defines)'}  ${m.importe:,.2f}")
            print(f"           {m.concepto}")
            exito = False
            for intento in (1, 2):
                try:
                    page = bot_web.pagina_viva(ctx, page)
                    dlg = bot_web.asegurar_dialogo(page)
                    if m.forma_pago.startswith("TARJETA_"):
                        etiqueta = resolver_forma_tarjeta(
                            page, dlg, m.forma_pago == "TARJETA_CREDITO")
                        if etiqueta:
                            m.forma_pago = etiqueta
                        else:
                            m.accion = "PENDIENTE"
                            m.nota = "no encontre opcion de tarjeta: eligela tu"
                            m.forma_pago = ""
                    seguro = m.accion == "CAPTURAR" and m.contacto and m.cat1
                    listo = bot_web.llenar_campos(page, dlg, m, auto=bool(seguro))
                    bot_web.captura(page, f"rep_{m.n}")

                    guardar_solo = seguro and listo
                    if guardar_solo:
                        est = bot_web.leer_estado(dlg)
                        for campo, valor in (("contacto", m.contacto), ("cat1", m.cat1),
                                             ("cat2", m.cat2), ("cat3", m.cat3)):
                            v_form, v_esp = bot_web._norm(est[campo]), bot_web._norm(valor or "")
                            if v_esp and not (v_esp in v_form or v_form in v_esp):
                                print(f"    (!) {campo}: quedo '{est[campo]}' y esperaba '{valor}'")
                                guardar_solo = False
                            if not v_esp and campo != "contacto" and v_form:
                                print(f"    (!) {campo}: quedo '{est[campo]}' (sobrante)")
                                guardar_solo = False
                    if guardar_solo:
                        bot_web.resolver_avisos_previos(dlg, page)
                        dlg.get_by_role("button",
                                        name=re.compile(r"Guardar y agregar otro", re.I)).first.click()
                        page.wait_for_timeout(900)
                        try:
                            page.get_by_role("button",
                                             name=re.compile(r"^S[ií]$|Confirmar|Aceptar", re.I)
                                             ).first.click(timeout=1500)
                            page.wait_for_timeout(400)
                        except PWTimeout:
                            pass
                        print("   ✔ guardado solo")
                    else:
                        print(f"   >>> Este lo confirmas TU "
                              f"({m.nota or 'no pude llenarlo completo'})")
                        bot_web.esperar_y_aprender(page, dlg, m, aprender=True)
                    ok += 1
                    exito = True
                    reg["web"][it["clave"]] = True
                    guardar_registro(reg)
                    break
                except Exception as e:
                    print(f"   (!) intento {intento} fallo: {str(e)[:90]}")
                    bot_web.captura(page, f"ERROR_rep_{m.n}")
                    page.wait_for_timeout(2500)
            if not exito:
                err += 1
                print(f"   ERROR (reintenta con: python bot_reporte.py --web --solo {m.n})")
        print(f"\n  ===== FIN WEB. Procesados: {ok}   Errores: {err} =====")
        try:
            page.wait_for_timeout(1500)
        except Exception:
            pass
        ctx.close()


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ingresos", nargs="?", help="xlsx de INGRESOS (def: el mas nuevo)")
    ap.add_argument("gastos", nargs="?", help="xlsx de GASTOS (def: el mas nuevo)")
    ap.add_argument("--excel", action="store_true", help="escribir en el Excel")
    ap.add_argument("--web", action="store_true", help="capturar en la web")
    ap.add_argument("--desde", type=int, default=1)
    ap.add_argument("--solo", type=int)
    args = ap.parse_args()

    ing = Path(args.ingresos) if args.ingresos else \
        max(DOCS.glob("INGRESOS DEL *.xlsx"), key=lambda p: p.stat().st_mtime)
    gas = Path(args.gastos) if args.gastos else \
        max(DOCS.glob("GASTOS DEL *.xlsx"), key=lambda p: p.stat().st_mtime)
    print(f"Reportes: {ing.name}  +  {gas.name}")

    provs, clientes, cnt, fechas = bot_excel.cargar_catalogos()
    detalles = detalles_existentes()
    items = construir(ing, gas, provs, detalles, fechas["GASTOS"])

    reg = cargar_registro()
    for it in items:
        k = it["clave"] = clave_item(it)
        dup_libro = bool(it["omitido"]) and ("ya esta" in it["omitido"]
                                             or "ya tiene" in it["omitido"])
        nuestro = k in reg["excel"]
        if nuestro and not it["omitido"]:
            it["omitido"] = "ya escrito en Excel (registro)"
        # web: se salta lo ya capturado en web y los duplicados PREVIOS a este
        # bot (gastos que venian de un corte PDF, capturados en web entonces);
        # lo que NOSOTROS escribimos al Excel si va a la web.
        if k in reg["web"]:
            it["omitido_web"] = "ya capturado en web (registro)"
        elif (dup_libro and not nuestro) or (
                it["omitido"] and "no cuadran" in it["omitido"]):
            it["omitido_web"] = it["omitido"]
        else:
            it["omitido_web"] = None

    print(f"\n{'mov':>4} {'fecha':6} {'hoja':8} {'concepto':58} {'pago':14} {'monto':>10}")
    tot_i = tot_g = 0
    for it in items:
        m = it["mov"]
        marca = "OMITIR: " + it["omitido"] if it["omitido"] else \
            ("web manual" if m.accion == "PENDIENTE" else "")
        pago = it["vals"][11] if it["hoja"] == "INGRESOS" and it["vals"] else \
            (it["vals"][8] if it["vals"] else m.forma_pago)
        print(f"{m.n:>4} {it['fecha']:%d/%m}  {it['hoja']:8} {m.concepto[:58]:58} "
              f"{str(pago):14} "
              f"${m.importe:>8,.2f}  {marca}")
        if not it["omitido"]:
            if it["hoja"] == "INGRESOS":
                tot_i += m.importe
            else:
                tot_g += m.importe
    n_cap = sum(1 for it in items if not it["omitido"])
    print(f"\n  {n_cap} filas a capturar | ingresos ${tot_i:,.2f} | gastos ${tot_g:,.2f}")

    if args.excel:
        pend = [it for it in items if not it["omitido"] and it["vals"]]
        filas = [(it["hoja"], it["vals"]) for it in pend]
        if not filas:
            print("\nEXCEL: nada nuevo que escribir.")
        else:
            print(f"\nEXCEL: preparando libro ({len(filas)} filas)...")
            bot_excel.correr_osascript(bot_excel.applescript_preparar())
            provs, clientes, cnt, fechas = bot_excel.cargar_catalogos()
            folio = {"INGRESOS": cnt["INGRESOS"], "GASTOS": cnt["GASTOS"]}
            for hoja, vals in filas:
                pref = "I" if hoja == "INGRESOS" else "G"
                vals[0] = f"{pref}{folio[hoja]:05d}"
                folio[hoja] += 1
            print("Escribiendo", len(filas), "filas...")
            bot_excel.correr_osascript(bot_excel.applescript_escribir(filas))
            for it in pend:
                reg["excel"][it["clave"]] = it["vals"][0]
            guardar_registro(reg)
            print(f"✔ Excel listo. Folios: I{cnt['INGRESOS']:05d}-I{folio['INGRESOS']-1:05d} "
                  f"y G{cnt['GASTOS']:05d}-G{folio['GASTOS']-1:05d}")

    if args.web:
        capturar_web(items, reg, desde=args.desde, solo=args.solo)

    if not args.excel and not args.web:
        print("\nENSAYO: corre con --excel y/o --web para capturar.")


if __name__ == "__main__":
    main()
